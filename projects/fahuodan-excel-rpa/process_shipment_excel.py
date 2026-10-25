# -*- coding: utf-8 -*-
"""
影刀可复用指令：发货单 Excel 取消合并单元格并规整为逐行完整数据。

发货单导出常把主字段做成纵向合并单元格，影刀按行读取时下方为空。
本指令会：取消合并 → 指定列向下填充 → 删空行 → 可选整行去重 → 另存。

在影刀中的用法：
1. 安装依赖：pip install openpyxl
2. 把本文件放进应用或自定义指令集。
3. 用「调用模块」选择 process_shipment_excel / process_excel / main。
4. 推荐 column_names：
   发货单号,物流中心编码,创建时间,发货单状态,物流商
5. sheet_name 为空时取第一个工作表；dedupe 默认 True。

样例数据见 data/examples/ 与 data/processed/。
"""

from __future__ import annotations

import json
import os
import re
from typing import Iterable, List, Optional, Sequence, Union

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# 与样例 Excel / 文档保持一致的推荐填充列
DEFAULT_FILL_COLUMNS = "发货单号,物流中心编码,创建时间,发货单状态,物流商"


def _to_bool(value, default=True):
    """影刀开关可能是字符串；bool(\"False\") 会变成 True，需单独解析。"""
    if value is None or value == "":
        return default
    if isinstance(value, str):
        text = value.strip().lower()
        if text in ("true", "1", "yes", "y", "是"):
            return True
        if text in ("false", "0", "no", "n", "否"):
            return False
        raise ValueError("dedupe 必须是 True/False，收到: %s" % value)
    return bool(value)


def _parse_column_names(column_names: Union[str, Sequence[str]]) -> List[str]:
    """将影刀传入的列名参数解析为列表。"""
    if column_names is None or column_names == "":
        raise ValueError("column_names 不能为空")
    if isinstance(column_names, (list, tuple)):
        names = [str(x).strip() for x in column_names]
    else:
        text = str(column_names).strip()
        if text.startswith("["):
            parsed = json.loads(text)
            if not isinstance(parsed, list):
                raise ValueError("column_names 的 JSON 必须是数组")
            names = [str(x).strip() for x in parsed]
        else:
            # 兼容逗号、中文逗号、分号、竖线、换行
            names = re.split(r"[,，;；|\n\r\t]+", text)
            names = [n.strip() for n in names]
    result = [n for n in names if n]
    if not result:
        raise ValueError("column_names 不能为空")
    return result


def _normalize_output_name(output_name: str) -> str:
    name = str(output_name or "").strip().strip('"').strip("'")
    if not name:
        raise ValueError("output_name 不能为空")
    if not name.lower().endswith((".xlsx", ".xlsm")):
        name = "%s.xlsx" % name
    return name


def _resolve_output_path(file_path: str, output_name: str) -> str:
    """
    解析输出路径：
    - 若 output_name 为绝对路径，或包含目录分隔符，则按其路径写入；
    - 否则写到源文件同目录下。
    """
    output_name = _normalize_output_name(output_name)
    if os.path.isabs(output_name) or os.path.dirname(output_name):
        output_path = os.path.abspath(output_name)
    else:
        output_path = os.path.join(os.path.dirname(file_path), output_name)
    out_dir = os.path.dirname(output_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    return output_path


def _header_map(ws: Worksheet, header_row: int = 1) -> dict:
    mapping = {}
    for col in range(1, ws.max_column + 1):
        value = ws.cell(header_row, col).value
        if value is None:
            continue
        key = str(value).strip()
        if key and key not in mapping:
            mapping[key] = col
    return mapping


def _unmerge_and_fill(ws: Worksheet) -> None:
    """取消全部合并单元格，并把原值写回合并区域内每一个单元格。"""
    merged_ranges = list(ws.merged_cells.ranges)
    for merged in merged_ranges:
        min_row, min_col = merged.min_row, merged.min_col
        max_row, max_col = merged.max_row, merged.max_col
        value = ws.cell(min_row, min_col).value
        ws.unmerge_cells(str(merged))
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                ws.cell(row, col).value = value


def _is_blank(value) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and value.strip() == "":
        return True
    return False


def _forward_fill_columns(
    ws: Worksheet, cols: Iterable[int], header_row: int = 1
) -> None:
    """对指定列做向下填充（空单元格沿用上一行有效值）。"""
    last_values = {col: None for col in cols}
    for row in range(header_row + 1, ws.max_row + 1):
        for col in cols:
            cell = ws.cell(row, col)
            if _is_blank(cell.value):
                if last_values[col] is not None:
                    cell.value = last_values[col]
            else:
                last_values[col] = cell.value


def _row_values(ws: Worksheet, row: int, max_col: int) -> List:
    return [ws.cell(row, col).value for col in range(1, max_col + 1)]


def _row_all_blank(values: Sequence) -> bool:
    return all(_is_blank(v) for v in values)


def _dedupe_preserve_order(rows: List[List]) -> List[List]:
    """按整行内容去重，保留首次出现顺序。"""
    seen = set()
    result = []
    for row in rows:
        key = tuple("" if _is_blank(v) else str(v).strip() for v in row)
        if key in seen:
            continue
        seen.add(key)
        result.append(row)
    return result


def _copy_column_widths(source: Worksheet, target: Worksheet) -> None:
    for col_idx in range(1, source.max_column + 1):
        letter = get_column_letter(col_idx)
        width = source.column_dimensions[letter].width
        if width:
            target.column_dimensions[letter].width = width


def _write_clean_workbook(
    headers: Sequence,
    rows: Sequence[Sequence],
    sheet_title: str = "发货单详情",
    source_ws: Optional[Worksheet] = None,
) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title or "发货单详情"

    for col, header in enumerate(headers, start=1):
        ws.cell(1, col).value = header

    for r_idx, row in enumerate(rows, start=2):
        for c_idx, value in enumerate(row, start=1):
            ws.cell(r_idx, c_idx).value = value

    if source_ws is not None:
        _copy_column_widths(source_ws, ws)

    # 明确保证没有任何合并单元格
    assert len(list(ws.merged_cells.ranges)) == 0
    return wb


def process_excel(
    file_path: str,
    column_names: Union[str, Sequence[str]],
    output_name: str,
    sheet_name: Optional[str] = None,
    dedupe: bool = True,
) -> str:
    """
    发货单 Excel：取消合并、指定列向下填充、可选整行去重后另存
    * @param file_path, 源文件路径（xlsx）
    * @param column_names, 需要向下填充的列字段名称，多个用逗号分隔或传列表
    * @param output_name, 处理后的文件名称或路径（可省略扩展名，默认 .xlsx）
    * @param sheet_name, 工作表名，默认取第一个
    * @param dedupe, 是否按整行去重（默认 True，适合发货单重复明细行）
    * @return `result`, 处理后文件的完整路径
    """
    file_path = os.path.abspath(str(file_path or "").strip().strip('"').strip("'"))
    if not file_path:
        raise ValueError("file_path 不能为空")
    if not os.path.isfile(file_path):
        raise FileNotFoundError("找不到文件: %s" % file_path)

    fill_names = _parse_column_names(column_names)
    do_dedupe = _to_bool(dedupe, default=True)
    output_path = _resolve_output_path(file_path, output_name)

    wb = load_workbook(file_path)
    try:
        if sheet_name is not None and str(sheet_name).strip() != "":
            name = str(sheet_name).strip()
            if name not in wb.sheetnames:
                raise ValueError(
                    "工作表不存在: %s，可用: %s" % (name, wb.sheetnames)
                )
            ws = wb[name]
        else:
            ws = wb[wb.sheetnames[0]]

        headers_map = _header_map(ws, header_row=1)
        if not headers_map:
            raise ValueError("未读取到表头，请确认第一行为字段名称")

        missing = [n for n in fill_names if n not in headers_map]
        if missing:
            raise ValueError(
                "以下列名在表头中不存在: %s；实际表头: %s"
                % (missing, list(headers_map.keys()))
            )

        fill_cols = [headers_map[n] for n in fill_names]
        max_col = ws.max_column
        headers = [ws.cell(1, c).value for c in range(1, max_col + 1)]

        _unmerge_and_fill(ws)
        _forward_fill_columns(ws, fill_cols, header_row=1)

        data_rows: List[List] = []
        for row_idx in range(2, ws.max_row + 1):
            values = _row_values(ws, row_idx, max_col)
            if _row_all_blank(values):
                continue
            data_rows.append(values)

        if do_dedupe:
            data_rows = _dedupe_preserve_order(data_rows)

        out_wb = _write_clean_workbook(
            headers,
            data_rows,
            sheet_title=ws.title,
            source_ws=ws,
        )
        out_wb.save(output_path)
        out_wb.close()
    finally:
        wb.close()

    return output_path


# 与模块名对齐的别名，便于影刀「调用模块」按文件名查找
process_shipment_excel = process_excel


def main(
    file_path: str,
    column_names: str = DEFAULT_FILL_COLUMNS,
    output_name: str = "发货单_已处理.xlsx",
    sheet_name: Optional[str] = None,
    dedupe=True,
) -> str:
    """
    影刀指令主入口
    * @param file_path, 源 Excel 文件路径
    * @param column_names, 需要向下填充的列名，逗号分隔；默认五列主字段
    * @param output_name, 处理后的文件名称或路径
    * @param sheet_name, 工作表名，默认取第一个（样例为「发货单详情」）
    * @param dedupe, 是否按整行去重，默认 True；影刀可传 \"True\"/\"False\"
    * @return `result`, 处理后文件完整路径
    """
    return process_excel(
        file_path,
        column_names,
        output_name,
        sheet_name=sheet_name,
        dedupe=dedupe,
    )


def run_from_args(args: dict) -> str:
    """
    影刀可视化流程入口（dict 参数）。
    流程参数：file_path、column_names、output_name，可选 sheet_name、dedupe。
    执行后写回 args[\"result\"]。
    """
    result = process_excel(
        args.get("file_path"),
        args.get("column_names", DEFAULT_FILL_COLUMNS),
        args.get("output_name", "发货单_已处理.xlsx"),
        sheet_name=args.get("sheet_name"),
        dedupe=args.get("dedupe", True),
    )
    args["result"] = result
    return result


if __name__ == "__main__":
    import argparse
    import sys

    parser = argparse.ArgumentParser(
        description="发货单 Excel 取消合并并规整（影刀可复用）"
    )
    parser.add_argument("file_path", help="源 Excel 文件路径")
    parser.add_argument(
        "column_names",
        nargs="?",
        default=DEFAULT_FILL_COLUMNS,
        help="需要向下填充的列名，逗号分隔（默认五列主字段）",
    )
    parser.add_argument(
        "output_name",
        nargs="?",
        default="发货单_已处理.xlsx",
        help="处理后的文件名称或路径",
    )
    parser.add_argument("--sheet", default=None, help="工作表名称，默认第一个")
    parser.add_argument(
        "--no-dedupe",
        action="store_true",
        help="不去重，仅取消合并并填充",
    )
    cli = parser.parse_args()

    try:
        result = process_excel(
            cli.file_path,
            cli.column_names,
            cli.output_name,
            sheet_name=cli.sheet,
            dedupe=not cli.no_dedupe,
        )
    except Exception as exc:  # noqa: BLE001 - CLI 需把错误完整返回给调用方
        print(
            json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False),
            file=sys.stderr,
        )
        sys.exit(1)

    print(result)
