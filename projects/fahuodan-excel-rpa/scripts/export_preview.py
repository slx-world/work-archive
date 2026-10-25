# -*- coding: utf-8 -*-
"""
导出可在 GitHub 在线查看的预览文件。

GitHub 不能直接预览 .xlsx（二进制），因此额外生成：
- CSV：表格内容
- 源表示例还可生成 Markdown（含合并区域说明）
"""

import csv
import sys
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value)


def export_csv(xlsx_path: Path, csv_path: Optional[Path] = None) -> Path:
    src = Path(xlsx_path)
    dest = Path(csv_path) if csv_path else src.with_name(src.stem + "_预览.csv")

    wb = load_workbook(src, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    with dest.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        for row in rows:
            writer.writerow([_cell_text(v) for v in row])
    return dest.resolve()


def export_source_markdown(xlsx_path: Path, md_path: Optional[Path] = None) -> Path:
    """源表 Markdown 预览：空白格对应合并区，并列出合并范围。"""
    src = Path(xlsx_path)
    dest = Path(md_path) if md_path else src.with_name(src.stem + "_预览.md")

    wb = load_workbook(src)
    try:
        ws = wb.active
        title = ws.title
        merges = sorted(str(r) for r in ws.merged_cells.ranges)
        rows = list(
            ws.iter_rows(
                min_row=1,
                max_row=ws.max_row,
                max_col=ws.max_column,
                values_only=True,
            )
        )
    finally:
        wb.close()

    def md_cell(value) -> str:
        text = _cell_text(value).replace("|", "\\|")
        return text if text else " "

    lines = [
        "# 发货单示例（含合并单元格）在线预览",
        "",
        "> GitHub 无法直接预览 `.xlsx`。本页为源表示意：空白单元格对应纵向合并区域。",
        "",
        "- 工作表：`%s`" % title,
        "- 合并区域数：**%d**" % len(merges),
        "- 明细行数（不含表头）：**%d**" % (len(rows) - 1),
        "",
        "## 表内容（合并区下方显示为空）",
        "",
        "| " + " | ".join(md_cell(v) for v in rows[0]) + " |",
        "| " + " | ".join(["---"] * len(rows[0])) + " |",
    ]
    for row in rows[1:]:
        lines.append("| " + " | ".join(md_cell(v) for v in row) + " |")

    lines.extend(["", "## 合并区域列表", ""])
    for item in merges:
        lines.append("- `%s`" % item)

    lines.extend(
        [
            "",
            "处理后结果见 [`../processed/发货单-示例-已处理_预览.csv`]"
            "(../processed/发货单-示例-已处理_预览.csv)。",
            "",
        ]
    )
    dest.write_text("\n".join(lines), encoding="utf-8")
    return dest.resolve()


def export_all_samples(project_root: Optional[Path] = None) -> None:
    root = Path(project_root) if project_root else Path(__file__).resolve().parents[1]
    src = root / "data" / "examples" / "发货单-示例-含合并单元格.xlsx"
    out = root / "data" / "processed" / "发货单-示例-已处理.xlsx"

    print(export_csv(src))
    print(export_source_markdown(src))
    if out.is_file():
        print(export_csv(out))


if __name__ == "__main__":
    if len(sys.argv) == 1:
        export_all_samples()
    elif sys.argv[1] in ("--all", "all"):
        export_all_samples(Path(sys.argv[2]) if len(sys.argv) > 2 else None)
    elif len(sys.argv) < 2:
        print("用法: python export_preview.py [<xlsx> [csv]] | --all")
        sys.exit(1)
    else:
        print(export_csv(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None))
